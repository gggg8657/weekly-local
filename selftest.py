#!/usr/bin/env python3
"""LLM 없이 검증 (가짜 LLM): 주차 → 약어집(사내>시드>공개, 이력) → 약어 검출·첫 등장 풀이 → 외부활동·날짜·분량 규칙 → 원문 대조(지어낸 숫자·장소)
→ 개인 항목화(지난 계획 대조) → 제출 → 미제출 부서 → 취합(합치기·빠진 항목 복원·표시 OR) → 압축(표시 항목 보존) → HWPX(색·취소선·병합·자리표시자, kordoc validate)
→ DOCX(XML) → BBS 게시용 → HTTP(SSE) → ui.html.
WORKSPACE 는 임시 폴더로 바꿔 실데이터 폴더에 흔적을 남기지 않는다.   python3 selftest.py"""
import base64
import datetime
import io
import json
import os
import re
import shutil
import sys
import tempfile
import threading
import urllib.error
import urllib.request
import zipfile
import xml.etree.ElementTree as ET

TMP = tempfile.mkdtemp(prefix="weekly-selftest-")
os.environ["WORKSPACE"] = os.path.join(TMP, "weekly-local")
os.environ["AGENT_DATA"] = TMP
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import app  # noqa: E402
import render  # noqa: E402
import rules  # noqa: E402

CALLS = []


def fake(system, user, model=None, temperature=0.2, on_token=lambda t: None, json_mode=True):
    CALLS.append((system[:20], user))
    if "주간보고 메모를" in system:
        items = [{"project": "기본사업", "period": "10/5~10/8", "text": "i-SMR 노심 해석 MARS-KS 대리모델 학습 완료", "kind": "done", "cat": "goal", "core": True},
                 {"project": "기본사업", "period": "10.7", "text": "IAEA 기술회의 발표", "kind": "done", "cat": "perf", "ext": True, "place": "오스트리아 빈", "people": "김연구, 박지어냄"},
                 {"project": "AI융합사업", "text": "과기정통부 3분기 실적 자료 제출", "kind": "done", "cat": "goal", "msit": True},  # 메모에 없는 과제명 → 추정 확인
                 {"text": "인사 관련 협의", "kind": "done", "cat": "etc", "nobbs": True},  # 과제명 없음 → 질의
                 {"project": "기본사업", "period": "10/15", "text": "KINS 규제 대응 회의 예정 (@KINS 대전, 김연구)", "kind": "plan", "cat": "perf"},
                 {"project": "기본사업", "text": "시험 결과 정리 95% 달성", "kind": "done", "cat": "goal"},  # 95 는 메모에 없음 → 경고
                 {"project": "기본사업", "text": "ZQX-3 코드 검증 완료", "kind": "done", "cat": "goal"},  # 약어집에 없는 약어 → 작성자 질의
                 {"project": "기본사업", "text": "열수력 시험 결과 정리", "kind": "done", "cat": "goal", "children": [  # 세부 항목(트리) → depth
                     {"text": "압력 데이터 후처리", "kind": "plan", "cat": "etc"}, {"text": "온도 데이터 비교", "children": [{"text": "12번 채널 이상치 제거", "nobbs": True}]}]}]
        carry = [{"n": 1, "status": "완료", "evidence": "대리모델 학습 완료"}, {"n": 2, "status": "완료", "evidence": "메모에 없는 근거"}]
        out = json.dumps({"items": items, "carry": carry, "remarks": ["시험 장비 고장으로 일정 지연 우려"]}, ensure_ascii=False)
    elif "약어 풀이 도우미" in system:
        out = json.dumps({"abbrs": [{"abbr": "ZQX-3", "full": "Zeta Quench eXperiment 3", "ko": "", "field": "원자력", "desc": ""}]}, ensure_ascii=False)
    elif "취합 담당자" in system:
        ids = re.findall(r"^\[([0-9a-f]+)\] ", user, flags=re.M)
        lines = {i: re.search(rf"^\[{i}\] .*$", user, flags=re.M).group(0) for i in ids}
        dup = [i for i in ids if "ATLAS" in lines[i].split(" ‹")[0]]
        out_items = []
        if dup:
            out_items.append({"src": dup, "text": "ATLAS 장기냉각 시험 2회 수행", "kind": "done", "cat": "goal"})
        rest = [i for i in ids if i not in dup]
        for i in rest[:-1]:  # 마지막 하나는 일부러 빠뜨림 → 원문 복원돼야 함
            t = re.sub(r"^\[\w+\] ", "", re.sub(r" ‹.*›$", "", lines[i]))
            t = re.sub(r" \(@.*\)$", "", t)
            out_items.append({"src": [i], "text": t, "kind": "plan" if "·계획" in lines[i] else "done", "cat": "perf"})
        out = json.dumps({"items": out_items}, ensure_ascii=False)
    elif "편집자" in system:
        keys = re.findall(r"^\[(r\d+i\d+)\]", user, flags=re.M)
        line = lambda k: re.search(rf"^\[{k}\].*$", user, flags=re.M).group(0)
        core = [k for k in keys if "·핵심" in line(k)]
        etc = [k for k in keys if "·etc" in line(k)]
        perf = [k for k in keys if "·perf" in line(k) and k not in core]
        out = json.dumps({"items": [{"id": core[0], "text": "짧게 고친 핵심 999건"}] if core else [], "drop": core[:1] + etc[:1] + perf[:1]}, ensure_ascii=False)
    else:
        out = "{}"
    if "FORMAT_BREAK" in user and "반드시" not in user:
        out = "죄송합니다, 형식 없이 답합니다"
    for i in range(0, len(out), 9):
        on_token(out[i:i + 9])
    return out


app.llm = fake
EMIT = lambda ev: None
try:
    assert app.WS.startswith(TMP), "WORKSPACE 가 임시 폴더가 아님"

    # 1) 주차
    w = rules.week_of("2026-10-07")
    assert w["key"] == "2026-W41" and w["mon"] == "2026-10-05" and w["fri"] == "2026-10-09", w
    assert w["label"] == "2026. 10. 5.(월) ~ 10. 9.(금)" and rules.week_of("2026-W41")["mon"] == "2026-10-05"

    # 2) 약어집: 사내 > 시드(국내·원자력 + AI·컴퓨터 큐레이션, 동음이의는 문맥) > 공개(NUREG-0544·NIST CSRC, 후보만), 이력
    gl = app.GL
    st = gl.stats()
    assert st["public_by_src"]["NUREG-0544 Rev.4"] > 7000 and st["public_by_src"]["NIST CSRC Glossary"] > 3000 and st["seed"] >= 90, st
    e, c = gl.lookup("KAERI")
    assert e["ko"] == "한국원자력연구원" and "시드" in e["src"] and e["field"] == "기관·정책"
    e, c, s_ = gl.resolve("EBS")  # 공개 약어집에만 있으면 자동 확정하지 않고 후보로
    assert e is None and s_ == "public" and len(c) >= 2 and "NUREG" in c[0]["src"]
    e, c, s_ = gl.resolve("GPU")
    assert e["full"] == "Graphics Processing Unit" and e["field"] == "컴퓨터"  # 큐레이션이 공개 약어집(옛 뜻)보다 우선
    gl.put("EC", "eddy current", "와전류", editor="테스터")
    gl.put("EC", "European Commission", "유럽연합 집행위원회", editor="테스터2", field="기관·정책")
    e, _ = gl.lookup("EC")
    assert e["full"] == "European Commission" and e["src"] == "사내" and gl.user()["EC"]["history"][0]["full"] == "eddy current"
    gl.put("KAERI", "Korea Atomic Energy Research Institute", "한국원자력연구원(사내)", editor="t")
    assert gl.lookup("KAERI")[0]["ko"] == "한국원자력연구원(사내)"  # 사내가 시드보다 우선
    gl.delete("KAERI")
    assert gl.lookup("KAERI")[0]["ko"] == "한국원자력연구원"
    assert any(r["tier"] == "공개" for r in gl.search("loss-of-coolant")) and all(r["field"] == "AI" for r in gl.search("", field="AI"))
    # 동음이의: 문맥 낱말로 정하고, 못 정하면 질의
    assert gl.resolve("DT", "DT 기반 플랜트 실시간 모니터링")[0]["full"] == "Digital Twin"
    assert gl.resolve("DT", "DT 분류 모델 학습")[0]["full"] == "Decision Tree"
    assert gl.resolve("DT", "DT 검토")[2] == "homonym"
    assert gl.resolve("PIE", "핵연료 PIE 수행")[0]["ko"] == "조사후시험" and gl.resolve("PIE", "PIE 빌드 옵션")[0]["field"] == "컴퓨터"
    assert gl.resolve("RL", "RL 검토")[2] == "homonym" and gl.resolve("RL", "RL 에이전트 보상 설계")[0]["ko"] == "강화학습"

    # 3) 약어 검출·풀이 수준
    toks = [t for t, _, _ in rules.find_abbrs("i-SMR 과 MARS-KS, APR1400 및 OECD/NEA 회의, Python 은 아님, 10 MW, KAERI의")]
    assert toks == ["i-SMR", "MARS-KS", "APR1400", "OECD/NEA", "MW", "KAERI"], toks
    out, notes, descs, unk, _ = rules.expand_items(["MARS-KS 계산 수행", "MARS-KS 재계산, 원자력안전위원회(NSSC) 보고", "10 MW 출력"], gl, mode="paren")
    assert out[0] == "MARS-KS(Multi-dimensional Analysis of Reactor Safety-KINS Standard) 계산 수행", out
    assert out[1] == "MARS-KS 재계산, 원자력안전위원회(NSSC) 보고" and out[2] == "10 MW 출력", out  # 두 번째 등장·이미 풀린 것·단위는 그대로
    x = lambda t, **k: rules.expand_items([t], gl, **{"mode": "paren", **k})[0][0]
    # 기본(원자력 전문가): 원자력 1 · 기관 2 · AI 3 · 컴퓨터 2
    assert x("SMR 설계") == "SMR(Small Modular Reactor) 설계"
    assert x("KINS 협의") == "KINS(Korea Institute of Nuclear Safety, 한국원자력안전기술원) 협의"
    assert x("RAG 구축") == "RAG(Retrieval-Augmented Generation, 검색 증강 생성: AI가 답변하기 전에 사내 문서나 데이터에서 관련 내용을 찾아 참고하게 하는 방식. 최신성과 근거성을 높이고 환각을 줄이는 데 목적이 있음.) 구축"
    assert x("GPU 증설") == "GPU(Graphics Processing Unit, 그래픽 처리 장치) 증설"
    assert x("SMR 설계", levels={"원자력": 2}) == "SMR(Small Modular Reactor, 소형모듈원자로) 설계"
    assert x("SMR 설계", levels=rules.PRESETS["일반 공개"]).startswith("SMR(Small Modular Reactor, 소형모듈원자로: ")
    o, n, d, u, _ = rules.expand_items(["RAG 구축"], gl, mode="paren", desc_block=True)
    assert o == ["RAG(Retrieval-Augmented Generation, 검색 증강 생성) 구축"] and d == [("RAG", "Retrieval-Augmented Generation, 검색 증강 생성", "AI가 답변하기 전에 사내 문서나 데이터에서 관련 내용을 찾아 참고하게 하는 방식. 최신성과 근거성을 높이고 환각을 줄이는 데 목적이 있음.")]
    o, n, d, u, _ = rules.expand_items(["SMR 1호기", "SMR 2호기"], gl, mode="paren", first_only=False)
    assert o[1].startswith("SMR(Small")
    # 미확인 약어는 지어내지 않고 표시만
    o, n, d, u, _ = rules.expand_items(["ZQXJ 코드 검증", "ZQXJ 재검증"], gl, mode="paren")
    assert o == ["ZQXJ" + rules.UNKNOWN + " 코드 검증", "ZQXJ 재검증"] and u == {"ZQXJ"}
    assert rules.expand_items(["실 세미나(LLM 기반 사례 공유)", "LOCA 해석"], gl, mode="paren")[0][0] == "실 세미나(LLM(Large Language Model, 대규모 언어모델: 대량의 문서를 학습해 질문 응답, 요약, 문서 작성, 정보 정리 등을 수행하는 AI. 사내 지식검색·보고서 작성·업무지원에 활용 가능.) 기반 사례 공유)"
    # 기본 = 풀이 줄(lines): 본문은 그대로, 약어가 처음 나온 항목 아래 한 줄씩 '약어: 풀이', 수준별
    o, n, d, u, ln = rules.expand_items(["SMR 설계와 RAG 구축", "KINS 협의, SMR 재검토", "ZQXJ 검증"], gl)
    assert o == ["SMR 설계와 RAG 구축", "KINS 협의, SMR 재검토", "ZQXJ 검증"], o
    assert ln[0] == [("SMR", "Small Modular Reactor"), ("RAG", "Retrieval-Augmented Generation(검색 증강 생성) — AI가 답변하기 전에 사내 문서나 데이터에서 관련 내용을 찾아 참고하게 하는 방식. 최신성과 근거성을 높이고 환각을 줄이는 데 목적이 있음.")], ln
    assert ln[1] == [("KINS", "Korea Institute of Nuclear Safety(한국원자력안전기술원)")] and ln[2] == [("ZQXJ", rules.UNKNOWN)], ln
    assert rules.expand_items(["SMR 설계"], gl, levels={"원자력": 2})[4][0] == [("SMR", "Small Modular Reactor(소형모듈원자로)")]
    assert not rules.already_expanded("sLLM(gemma) 기반", 0, 4) and rules.already_expanded("RAG(Retrieval Augmented) 구축", 0, 3)
    assert not rules.already_expanded("세미나(LLM 기반)", 4, 7) and rules.already_expanded("원자력안전위원회(NSSC) 보고", 9, 13)
    o, n, d, u, _ = rules.expand_items(["KINS 회의"], gl, mode="note")
    assert o == ["KINS 회의"] and n == [("KINS", "Korea Institute of Nuclear Safety(한국원자력안전기술원)")]
    ws = rules.abbr_warnings([{"id": "x", "text": "ZQXJ 코드와 EBS 검토, KINS 협의, DT 검토"}], gl)
    st = {w["abbr"]: (w["status"], w["level"]) for w in ws}
    assert st == {"ZQXJ": ("missing", "ask"), "EBS": ("public", "ask"), "KINS": ("ok", "info"), "DT": ("homonym", "ask")}, st
    assert [c["full"] for c in next(w for w in ws if w["abbr"] == "DT")["candidates"]][:2] == ["Digital Twin", "Decision Tree"]
    gl.put("ZQND", "Zq No Desc", field="일반")  # 설명 없는 항목 → 수준 3 에서 'nodesc'
    assert rules.abbr_warnings([{"id": "y", "text": "ZQND 계획"}], gl, levels={"일반": 3})[0]["status"] == "nodesc"
    gl.delete("ZQND")

    # 4) 외부활동·날짜·분량 규칙
    ext = rules.ext_warnings([{"id": "a", "text": "IAEA 회의 발표"}, {"id": "b", "text": "학회 발표", "place": "제주", "people": "김"},
                              {"id": "c", "text": "실 내부 회의"}, {"id": "d", "text": "워크숍 참석 (@서울, 홍)"},
                              {"id": "e", "text": "한국원자력학회 우수논문상 수상", "ext": True}, {"id": "f", "text": "설계 검토 회의", "ext": False},
                              {"id": "g", "text": "설계 검토 회의"}])
    assert [w["item"] for w in ext] == ["a", "g"], ext
    assert rules.normalize_dates("10/15 회의, 2026-10-20 마감, 10월 22일(목) 발표, 10. 23. 출장, 2027-01-05", 2026) == \
        "10.15 회의, 10.20 마감, 10.22 발표, 10.23 출장, '27.1.5"  # 공식 양식 표기 M.D
    assert rules.normalize_dates("수율 3.5% 향상", 2026) == "수율 3.5% 향상"
    assert [rules.norm_period(x, 2026) for x in ("(2/24~3/5)", "~3월 6일", "3.20", "10. 14.(수) ~ 10. 16.(금)")] == ["2.24~3.5", "~3.6", "3.20", "10.14~10.16"]
    assert rules.period_ok("2.24~3.5") and rules.period_ok("~3.6") and not rules.period_ok("다음주")
    assert rules.date_warnings([{"text": "10/15 회의"}, {"text": "10월 16일 회의"}])

    # 5) 원문 대조
    it = {"text": "시험 3회 수행, 수율 95%", "place": "대전 KINS", "people": "김연구, 이없음"}
    wn = rules.verify_item(it, "시험 3회 수행 (@대전 KINS, 김연구)")
    assert any("95" in x for x in wn) and any("이없음" in x for x in wn) and it["people"] == "김연구", (wn, it)

    assert app.norm_item({"text": "(핵심)(goal) ATLAS 시험 2회 ‹열수력·수행·goal·핵심›"})["text"] == "ATLAS 시험 2회"  # LLM 이 참고 표시를 베낀 경우
    assert app.strip_dup_place("워크숍 준비(10월 22일, 본원 교육센터)", "본원 교육센터", "") == "워크숍 준비(10월 22일)"

    # 6) 지난 주 제출 → 이번 주 항목화(지난 계획 대조)
    P1 = {"name": "김연구", "org": "가상원자력연구소", "dept": "인공지능응용연구실"}
    app.submit({"profile": P1, "week": "2026-09-30", "items": [
        {"project": "기본사업", "text": "MARS-KS 대리모델 학습", "kind": "plan", "cat": "goal"}, {"project": "기본사업", "text": "특허 출원서 작성", "kind": "plan", "cat": "perf"}]})
    memo = "[기본사업]\n- ★ i-SMR 노심 해석 MARS-KS 대리모델 학습 완료(10/5~10/8)\n- IAEA 기술회의 발표 10/7 (오스트리아 빈, 김연구)\n- 과기정통부 보고 3분기 실적 자료 제출\n- 인사 관련 협의 (비공개)\n다음주: KINS 규제 대응 회의 10/15 (KINS 대전, 김연구)\n- 시험 결과 정리\n- ZQX-3 코드 검증 완료\n- 열수력 시험 결과 정리\n  - 압력 데이터 후처리\n  - 온도 데이터 비교\n    - 12번 채널 이상치 제거"
    r = app.run_itemize({"profile": P1, "week": "2026-10-07", "memo": memo}, EMIT, "fake")
    its = r["items"]
    assert len(its) == 11 and r["prev_week"] == "2026-W40", r["prev_week"]
    tree = [(i["text"], i["depth"], i["kind"], i["cat"]) for i in its[-4:]]
    assert tree == [("열수력 시험 결과 정리", 0, "done", "goal"), ("압력 데이터 후처리", 1, "done", "goal"), ("온도 데이터 비교", 1, "done", "goal"),
                    ("12번 채널 이상치 제거", 2, "done", "goal")], tree  # 하위는 상위의 수행/계획·분류를 따름
    q = [x for x in r["questions"] if x["kind"] == "abbr"]
    assert [x["abbr"] for x in q] == ["ZQX-3"] and q[0]["sentence"].startswith("ZQX-3 코드") and q[0]["level"] == "ask", q
    qp = {x["status"]: x for x in r["questions"] if x["kind"] == "project"}  # 과제명 없음 → 질의, 메모에 없는 과제명 → 확인
    assert set(qp) == {"missing", "guess"} and qp["guess"]["suggest"] == "AI융합사업" and "인사" in qp["missing"]["sentence"], r["questions"]
    assert r["remarks"] == ["시험 장비 고장으로 일정 지연 우려"]
    isr = next(i for i in its if "i-SMR" in i["text"])
    assert isr["project"] == "기본사업" and isr["period"] == "10.5~10.8" and not isr.get("project_guess"), isr
    sug = [c for c in q[0]["candidates"] if c.get("suggest")]
    assert sug and sug[0]["src"] == "LLM 제안(확인 필요)"  # LLM 제안은 후보로만
    assert "ZQX-3: " + rules.UNKNOWN in render.to_text(app.personal_doc({"week": "2026-W41", "dept": "d", "name": "n", "items": its}), gl, render.template())
    assert "Zeta" not in render.to_text(app.personal_doc({"week": "2026-W41", "dept": "d", "name": "n", "items": its}), gl, render.template())
    iaea = next(i for i in its if "IAEA" in i["text"])
    assert iaea["people"] == "김연구" and any("박지어냄" in w for w in iaea["warn"]), iaea  # 메모에 없는 참석자 지움
    kins = next(i for i in its if "KINS" in i["text"])
    assert kins["text"] == "KINS 규제 대응 회의 예정" and kins["period"] == "10.15" and kins["place"] == "KINS 대전" and kins["people"] == "김연구", kins  # (@…)·기간을 칸으로
    assert any("95" in w for w in next(i for i in its if "95%" in i["text"])["warn"])
    assert [c["status"] for c in r["carry"]] == ["완료", "미확인"], r["carry"]  # 메모에 없는 근거의 '완료' 는 미확인으로
    it_call = [u for sy, u in CALLS if sy.startswith("너는 한국원자력연구원(KAERI)")][-1]
    assert "[지난 계획]" in it_call and "1. MARS-KS 대리모델 학습" in it_call

    # 형식이 깨지면 한 번 더
    n0 = len(CALLS)
    app.run_itemize({"profile": P1, "week": "2026-10-07", "memo": "FORMAT_BREAK 시험 수행", "use_prev": False}, EMIT, "fake")
    assert len([1 for sy, _ in CALLS[n0:] if sy.startswith("너는 한국원자력연구원(KAERI)")]) == 2

    # 7) 제출 → 미제출 부서
    app.save_orgs({"가상원자력연구소": ["인공지능응용연구실", "열수력안전연구실", "원자로설계실"]})
    items1 = [i for i in its if "95%" not in i["text"]]
    try:  # 미확인 약어가 있으면 제출 차단
        app.submit({"profile": P1, "week": "2026-10-07", "items": items1, "memo": memo})
        raise AssertionError("미확인 약어가 있는데 제출됨")
    except app.Unresolved as e:
        assert [w.get("abbr") for w in e.rows if w["kind"] == "abbr"] == ["ZQX-3"] and len([w for w in e.rows if w["kind"] == "project"]) == 2
    for i in items1:  # 작성자가 과제명을 정함(질의 카드 답)
        if not i.get("project"):
            i["project"] = "기본사업"
        i.pop("project_guess", None)
    # 작성자 답변 → 사내 약어집 저장(이력) → 재사용
    gl.put("ZQX-3", "Zeta Quench eXperiment 3", "제타 급랭 실험 3", editor="김연구", src="작성자 답변", field="원자력", ctx="ZQX-3 코드 검증 완료")
    assert gl.user()["ZQX-3"]["src"] == "작성자 답변" and gl.user()["ZQX-3"]["editor"] == "김연구"
    app.submit({"profile": P1, "week": "2026-10-07", "items": items1, "memo": memo})
    r2 = app.run_itemize({"profile": P1, "week": "2026-10-07", "memo": memo}, EMIT, "fake")
    assert not [x for x in r2["questions"] if x["kind"] == "abbr"] and "ZQX-3: Zeta Quench eXperiment 3" in render.to_text(app.personal_doc({"week": "2026-W41", "dept": "d", "items": r2["items"]}), gl, render.template())
    # 사유를 적으면 미확인 약어가 있어도 제출(사유 기록)
    rr = app.submit({"profile": {"name": "최사유", "org": "다른연구소", "dept": "x"}, "week": "2026-10-07",
                     "items": [{"text": "QQZW 장비 점검", "kind": "done", "cat": "etc"}], "unresolved_reason": "장비 고유 명칭, 풀이 없음"})
    saved = json.load(open(os.path.join(app.WS, rr["path"])))
    assert saved["unresolved"][0]["abbr"] == "QQZW" and saved["unresolved_reason"]
    app.submit({"profile": {"name": "이열수", "org": "가상원자력연구소", "dept": "열수력안전연구실"}, "week": "2026-10-07", "items": [
        {"project": "ATLAS 과제", "period": "10.6~10.8", "text": "ATLAS 장기냉각 시험 2회 수행", "kind": "done", "cat": "goal", "core": True},
        {"text": "NURETH 학회 발표", "kind": "done", "cat": "perf", "ext": True, "place": "부산 BEXCO", "people": "이열수"},
        {"text": "시험 보고서 작성", "kind": "plan", "cat": "etc"}], "unresolved_reason": "학회·장소 이름 — 취합 때 확인"})
    app.submit({"profile": {"name": "박시험", "org": "가상원자력연구소", "dept": "열수력안전연구실"}, "week": "2026-10-07", "items": [
        {"project": "ATLAS 과제", "period": "10.6~10.8", "text": "ATLAS 시험 2회 수행 지원", "kind": "done", "cat": "goal", "msit": True},
        {"project": "ATLAS 과제", "period": "10.9", "text": "계측기 교정", "kind": "done", "cat": "etc"}]})
    st = app.status("2026-W41", "가상원자력연구소")
    miss = [d["dept"] for d in st["depts"] if d["missing"]]
    assert miss == ["원자로설계실"], st
    try:
        app.submit({"profile": {"name": "", "dept": "x"}, "items": [{"text": "a"}]})
        raise AssertionError("이름 없는 제출 통과")
    except ValueError:
        pass

    # 8) 취합: 합치기 + 표시 OR + 장소 모으기 + 빠진 항목 복원
    agg = app.run_aggregate({"week": "2026-10-07", "org": "가상원자력연구소"}, EMIT, "fake")
    rows = {r["label"]: r for r in agg["rows"]}
    assert rows["원자로설계실"].get("missing") and not rows["원자로설계실"]["items"]
    th = rows["열수력안전연구실"]["items"]
    atlas = next(i for i in th if "ATLAS" in i["text"])
    assert atlas["core"] and atlas["msit"] and len(atlas["src"]) == 2 and atlas["project"] == "ATLAS 과제" and atlas["period"] == "10.6~10.8", atlas
    restored = [i for i in th if "원문 그대로 복원" in " ".join(i.get("warn") or [])]
    assert len(restored) == 1, th
    ai = rows["인공지능응용연구실"]["items"]
    assert any("복원" in " ".join(i.get("warn") or []) for i in ai) and len(ai) == len(items1)
    assert os.path.exists(app.agg_path("2026-W41", "가상원자력연구소"))
    seq = [(i["text"], i["depth"]) for i in ai]
    k = seq.index(("온도 데이터 비교", 1))
    assert seq[k + 1] == ("12번 채널 이상치 제거", 2) and [u for _, u in CALLS if u.startswith("[부서]")] and not any("12번" in u for _, u in CALLS if u.startswith("[부서]")), seq  # 하위는 LLM 에 안 보내고 그대로 붙임
    asks = [w for w in rules.check_doc(app.agg_doc(agg), gl) if w.get("level") == "ask"]
    nu = next(w for w in asks if w.get("abbr") == "NURETH")  # 취합자가 누구에게 물을지
    assert nu["dept"] == "열수력안전연구실" and nu["who"] == ["이열수"], nu

    # 9) 1쪽 추정·압축 (표시 항목은 지우지 않음)
    tpl = render.template()
    HAS_FORM = os.path.exists(os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates", "kaeri_weekly.hwpx"))
    big = json.loads(json.dumps(agg))
    for r_ in big["rows"]:
        if r_["items"]:
            r_["items"] += [dict(r_["items"][-1], id=f"z{k}", text="추가 업무 항목 " * 6, cat="etc", core=False, msit=False) for k in range(12)]
    est = rules.page_estimate(app.agg_doc(big), tpl["page"])
    assert est["over"], est
    assert any(w["kind"] == "page" for w in rules.check_doc(app.agg_doc(big), gl, tpl["page"]))
    res = app.run_compress({"agg": big}, EMIT, "fake")
    flat = [i for r_ in res["agg"]["rows"] for i in r_["items"]]
    assert any(i["core"] for i in flat) and any("표시 항목은 지우지 않음" in n for n in res["notes"]), res["notes"]
    assert any("'기타 업무' 항목이 남아 있어" in n for n in res["notes"]), res["notes"]  # (b) 항목은 기타가 남아 있는 동안 못 지움
    assert not any("999" in i["text"] for i in flat)  # 원문에 없는 숫자로 고친 것은 반영 안 함
    assert sum(len(r_.get("dropped") or []) for r_ in res["agg"]["rows"]) == 1

    kc = None
    doc = app.agg_doc(agg)
    pd = app.personal_doc({"week": "2026-W41", "org": "가상원자력연구소", "dept": "인공지능응용연구실", "name": "김연구", "items": items1, "remarks": ["장비 고장"]})
    if not HAS_FORM:
        print("  (건너뜀) 기관 공식 양식(templates/kaeri_weekly.*)이 없어 10~12번 양식 복제 검사를 건너뜁니다")
    else:
        # 10) HWPX — 공식 양식(templates/kaeri_weekly.hwpx) 복제: 표 2개(수행·계획), 실 행 복제, ∙ (과제명) 굵게 → - (기간) 내용, 색·취소선, 분류 제목 없음
        assert tpl["engine"] == "form" and render.template("default")["hwpx"] == "default.hwpx"
        doc = app.agg_doc(agg)
        hw = render.to_hwpx(doc, gl, tpl)
        ins = render.inspect_hwpx(hw)
        assert ins["colors"].get("#FF6600") and ins["colors"].get("#0000FF") and ins["strike_runs"] == 2, ins["colors"]
        assert ins["tables"] == 2 and ins["rows"] == 2 * (1 + 3) and ins["merged"].count((2, 1)) == 2 * (1 + 3), ins["merged"]  # 머리행 + 실 3행, 내용 칸 2열 병합(양식대로)
        txt = " ".join(t for t, _, _ in ins["texts"])
        assert "주간업무보고(서면보고용)" in txt and "Ⅰ. 가상원자력연구소" in txt and "‘26.10.5.~’26.10.9." in txt and "‘26.10.12.~’26.10.23." in txt and "열수력안전연구실" in txt, txt[:300]
        assert "(기본사업)" in ins["bold_texts"] and "(ATLAS 과제)" in ins["bold_texts"] and txt.count("(기본사업)") == 2, txt  # 과제명 한 번씩(수행·계획)
        assert "중점목표" not in txt and "기타 업무" not in txt and "※ 핵심사항" not in txt  # 분류 제목·작성 지침 없음
        assert "(10.5~10.8) " in txt and "(@부산 BEXCO, 이열수)" in txt and "인사 관련 협의" in txt and "3. 특기 및 애로사항" in txt
        assert "※ 약어" in txt and "i-SMR" in ins["bold_texts"] and "BEXCO" in ins["bold_texts"] and ": innovative Small Modular Reactor" in txt  # 한 줄에 하나, 약어 굵게
        strike_t = [t for t, _, s in ins["texts"] if s]
        assert any("인사 관련 협의" in t for t in strike_t) and any("12번 채널" in t for t in strike_t)
        sec = zipfile.ZipFile(io.BytesIO(hw)).read("Contents/section0.xml").decode()
        assert 'pageBreak="1"' not in sec and "가상원자로연구실" not in sec and "전략개발단" not in sec  # 양식 견본 글은 남지 않음
        bbs = render.inspect_hwpx(render.to_hwpx(doc, gl, tpl, bbs=True))
        assert bbs["strike_runs"] == 0 and "인사 관련 협의" not in " ".join(t for t, _, _ in bbs["texts"])
        kc = app.kordoc_check(hw)
        if kc is not None:
            assert kc["valid"], kc["msg"]
            assert "ATLAS" in kc["text"] and "수행업무" in kc["text"] and "향후 2주 계획" in kc["text"], kc["text"][:500]
        keep = render.inspect_hwpx(render.to_hwpx(dict(doc, keep_guide=True), gl, tpl))
        assert "※ 핵심사항" in " ".join(t for t, _, _ in keep["texts"])  # 옵션으로 작성 지침 유지
        # 기본(임시) 양식도 같은 구조
        dins = render.inspect_hwpx(render.to_hwpx(doc, gl, render.template("default")))
        assert dins["tables"] == 2 and "(기본사업)" in dins["bold_texts"] and "중점목표" not in " ".join(t for t, _, _ in dins["texts"])
        pd = app.personal_doc({"week": "2026-W41", "org": "가상원자력연구소", "dept": "인공지능응용연구실", "name": "김연구", "items": items1, "remarks": ["장비 고장"]})
        # 들여쓰기: 문단 속성(왼쪽 여백·내어쓰기)으로 — 공백 흉내 아님. ∙ 과제명 → - 항목 → · 하위
        pins = render.inspect_hwpx(render.to_hwpx(pd, gl, tpl))
        ind = {}
        for t, l, i in pins["indents"]:
            for b in ("∙ ", "- ", "· "):
                if t.startswith(b):
                    ind.setdefault(b, (l, i))
        em11 = 1100
        assert ind["∙ "] == (0, -round(rules.em("∙ ") * em11)), ind
        assert ind["- "] == (round(0.5 * em11), -round(1.2 * em11)) and ind["· "][0] == round(2.0 * em11), ind  # 0.5em, 하위는 +1.5em 안으로
        assert not any(re.match(r"^\s+(?:[·∙]|-\s*\()", t) for t, _, _ in pins["texts"])  # 글머리 앞 공백 없음(양식 견본의 ' - ' 대신 문단 여백)
        assert " - 장비 고장" in " ".join(t for t, _, _ in pins["texts"])  # 특기사항
        dx2 = zipfile.ZipFile(io.BytesIO(render.to_docx(pd, gl, tpl))).read("word/document.xml").decode()
        assert '<w:ind w:left="374" w:hanging="264"/>' in dx2, re.findall(r"<w:ind [^>]*/>", dx2)[:8]  # 11pt: 0.5em(110)+1.2em(264)
        assert "12번 채널" not in render.to_text(pd, gl, tpl, bbs=True) and "온도 데이터 비교" in render.to_text(pd, gl, tpl, bbs=True)
        assert "(김연구)" in " ".join(t for t, _, _ in pins["texts"])

        # 11) DOCX
        dx = render.to_docx(doc, gl, tpl)
        z = zipfile.ZipFile(io.BytesIO(dx))
        d = z.read("word/document.xml").decode()
        ET.fromstring(d)
        assert "<w:strike/>" in d and 'w:val="FF6600"' in d and 'w:val="0000FF"' in d and d.count("<w:tbl>") == 2 and "중점목표" not in d
        assert '<w:b/></w:rPr><w:t xml:space="preserve">(기본사업)</w:t>' in d and "주간업무보고(서면보고용)" in d
        assert "<w:strike/>" not in zipfile.ZipFile(io.BytesIO(render.to_docx(doc, gl, tpl, bbs=True))).read("word/document.xml").decode()

        # 12) 텍스트·HTML
        t = render.to_text(doc, gl, tpl)
        assert "[핵심]" in t and "[비게시]" in t and "■ 열수력안전연구실" in t and "∙ (기본사업)" in t and "- (10.5~10.8) " in t
        assert "인사 관련" not in render.to_text(doc, gl, tpl, bbs=True)
        h = render.to_html(doc, gl, tpl, full=True)
        assert "line-through" in h and "#FF6600" in h and h.count('<table class="wk-table"') == 2 and 'data-kind="plan"' in h and "중점목표" not in h

    # 12-1) 가져오기·내보내기 — 왕복(넣어 둔 JSON), 엑셀 입력 양식, 손으로 쓴 공식 양식(HWPX·DOCX), 지저분한 엑셀, 형식 섞인 여러 파일
    import exchange
    import xlsx
    keyset = ("text", "project", "period", "depth", "kind", "core", "msit", "nobbs", "place", "people")
    norm = lambda its: sorted(json.dumps({k: (i.get(k) or ("" if k in ("text", "project", "period", "place", "people", "kind") else 0 if k == "depth" else False)) for k in keyset}, ensure_ascii=False) for i in its)
    want = {r_["label"]: norm(r_["items"]) for r_ in agg["rows"] if r_["items"]}
    for fmt in ("hwpx", "docx", "xlsx", "xlsx-data"):
        data, _, name = app.export({"agg": agg, "format": fmt})
        res = app.import_files({"files": [{"name": name, "b64": base64.b64encode(data).decode()}], "week": "2026-W41"})["files"][0]
        got = {}
        for r_ in res["rows"]:
            got.setdefault(r_["dept"], []).append(r_["item"])
        assert res["method"] == "embedded" and {k: norm(v) for k, v in got.items()} == want, (fmt, res["method"])  # 손실 없는 왕복
    hw = app.export({"agg": agg, "format": "hwpx"})[0]
    assert app.kordoc_check(hw) is None or app.kordoc_check(hw)["valid"]  # 넣어 둔 JSON 이 있어도 구조 검증 통과
    bbs_hw = app.export({"agg": agg, "format": "hwpx", "bbs": True})[0]
    assert "인사 관련" not in exchange.find_embedded("x.hwpx", bbs_hw).__repr__()  # 게시용에는 비게시 항목을 넣지 않음

    def strip_embed(data):
        zin, buf = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            for i in zin.infolist():
                if "weekly-local" not in i.filename and "customXml" not in i.filename:
                    z.writestr(i, zin.read(i))
        return buf.getvalue()
    for fmt in ("hwpx", "docx"):  # JSON 없이 보고서 모양만으로(손으로 쓴 것과 같은 경로)
        res = exchange.import_file("x." + fmt, strip_embed(app.export({"agg": agg, "format": fmt})[0]), "2026-W41")
        got = {}
        for r_ in res["rows"]:
            got.setdefault(r_["dept"], []).append(app.norm_item(r_["item"]))
        cmp_ = lambda its: norm([{k: v for k, v in i.items() if k not in ("place", "people")} | {"text": i["text"] + rules.ext_suffix(i)} for i in its])
        A, B = {k: cmp_(v) for k, v in got.items()}, {r_["label"]: cmp_(r_["items"]) for r_ in agg["rows"] if r_["items"]}
        assert res["method"] == "form" and A == B, (fmt, [(k, set(A.get(k, [])) ^ set(B.get(k, []))) for k in B if A.get(k) != B.get(k)])
        assert (res["org"] == "가상원자력연구소" or not HAS_FORM) and any(a["abbr"] == "i-SMR" for a in res["abbrs"]), (res["org"], res["abbrs"][:3])
    if not HAS_FORM:
        print("  (건너뜀) 기관 공식 양식이 없어 format.hwpx 빈 양식·손으로 채운 양식 가져오기 검사를 건너뜁니다")
    else:
        # 공식 양식(format.hwpx) — 빈 양식 그대로, 그리고 양식에 손으로 채운 것
        blank = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "templates", "kaeri_weekly.hwpx"), "rb").read()
        res = exchange.import_file("format.hwpx", blank, "2026-W10")
        assert res["method"] == "form" and res["org"] == "선진원자로연구소" and not res["rows"]
        zin, buf = zipfile.ZipFile(io.BytesIO(blank)), io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            for i in zin.infolist():
                b = zin.read(i)
                if i.filename == "Contents/section0.xml":
                    x = b.decode()
                    x = x.replace("<hp:t>(2.24~3.5) </hp:t>", "<hp:t>(2.24~3.5) 노심 해석 코드 검증 완료</hp:t>", 1)
                    x = x.replace('<hp:run charPrIDRef="24"><hp:t> - (~3.6) </hp:t>', '<hp:run charPrIDRef="26"><hp:t> - (~3.6) 과기정통부 실적 자료 제출</hp:t>', 1)
                    b = x.encode()
                z.writestr(i, b)
        res = exchange.import_file("손으로쓴.hwpx", buf.getvalue(), "2026-W10")
        got = {(r_["dept"], r_["item"]["project"], r_["item"]["period"], r_["item"]["text"], bool(r_["item"].get("msit"))) for r_ in res["rows"]}
        assert ("가상원자로연구실", "기본사업", "2.24~3.5", "노심 해석 코드 검증 완료", False) in got and \
            ("인공지능응용연구실", "전략개발단", "~3.6", "과기정통부 실적 자료 제출", True) in got, got  # 파랑 글자 → 과기정통부, 두 줄 실 이름 합침

    # 손으로 만든 DOCX(워드에서 표를 그리고 색·취소선을 준 모양): 머리 '1. 수행업무', 실 | 내용, 굵은 (과제명), 파랑·주황·취소선, 들여쓰기
    W_ = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    run = lambda t, pr="": f'<w:r><w:rPr>{pr}</w:rPr><w:t xml:space="preserve">{t}</w:t></w:r>'
    par = lambda runs, ind="": f'<w:p><w:pPr>{ind}</w:pPr>{runs}</w:p>'
    cell = lambda ps: f"<w:tc><w:tcPr/>{ps}</w:tc>"
    body = (par(run("Ⅰ. 손작성연구소", "<w:b/>")) + "<w:tbl><w:tblPr/>"
            + "<w:tr>" + cell(par(run("1. 수행업무", "<w:b/>"))) + cell(par(run("'26.10.5.~'26.10.9."))) + "</w:tr>"
            + "<w:tr>" + cell(par(run("디지털원자로실"))) + cell(par(run("(디지털트윈 과제)", "<w:b/>")) + par(run("- (10.5~10.7) ") + run("가상 운전 시험 완료", '<w:color w:val="ED7D31"/>'), '<w:ind w:left="200"/>')
                                                            + par(run("· 시나리오 12종 검증"), '<w:ind w:left="600"/>') + par(run("- ") + run("내부 인사 협의", "<w:strike/>"), '<w:ind w:left="200"/>')) + "</w:tr></w:tbl>"
            + "<w:tbl><w:tblPr/><w:tr>" + cell(par(run("2. 향후 2주 계획"))) + cell(par(run(""))) + "</w:tr><w:tr>" + cell(par(run("디지털원자로실")))
            + cell(par(run("∙ (디지털트윈 과제)")) + par(run("- (~10.20) ") + run("과기정통부 중간보고", '<w:color w:val="0070C0"/>'), '<w:ind w:left="200"/>')) + "</w:tr></w:tbl>"
            + par(run("3. 특기 및 애로사항")) + par(run("- 서버 이전 일정 확정 필요")))
    dbuf = io.BytesIO()
    with zipfile.ZipFile(dbuf, "w") as z:
        z.writestr("[Content_Types].xml", render.CT)
        z.writestr("_rels/.rels", render.RELS)
        z.writestr("word/document.xml", f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:document {W_}><w:body>{body}</w:body></w:document>')
    res = exchange.import_file("손작성.docx", dbuf.getvalue(), "2026-W41")
    got = [(r_["item"]["kind"], r_["item"]["project"], r_["item"]["period"], r_["item"]["text"], r_["item"]["depth"], bool(r_["item"].get("core")),
            bool(r_["item"].get("msit")), bool(r_["item"].get("nobbs"))) for r_ in res["rows"]]
    assert res["method"] == "form" and res["org"] == "손작성연구소" and res["remarks"] == ["서버 이전 일정 확정 필요"], res
    assert got == [("done", "디지털트윈 과제", "10.5~10.7", "가상 운전 시험 완료", 0, True, False, False),
                   ("done", "디지털트윈 과제", "", "시나리오 12종 검증", 1, False, False, False),
                   ("done", "디지털트윈 과제", "", "내부 인사 협의", 0, False, False, True),
                   ("plan", "디지털트윈 과제", "~10.20", "과기정통부 중간보고", 0, False, True, False)], got
    # 보고서형 엑셀을 JSON 없이(손으로 같은 모양으로 만든 것과 같은 경로)
    rx = app.export({"agg": agg, "format": "xlsx"})[0]
    zin, buf2 = zipfile.ZipFile(io.BytesIO(rx)), io.BytesIO()
    wbx = zin.read("xl/workbook.xml").decode()
    with zipfile.ZipFile(buf2, "w") as z:
        for i in zin.infolist():
            b_ = zin.read(i)
            if i.filename == "xl/workbook.xml":
                b_ = re.sub(r'<sheet name="_data"[^>]*/>', "", wbx).encode()
            z.writestr(i, b_)
    res = exchange.import_file("보고서형.xlsx", buf2.getvalue(), "2026-W41")
    assert res["method"] == "form" and len(res["rows"]) == sum(len(r_["items"]) for r_ in agg["rows"]), (res["method"], len(res["rows"]))
    assert any(r_["item"].get("core") and r_["item"].get("msit") for r_ in res["rows"]) and any(r_["item"].get("nobbs") for r_ in res["rows"])

    # 엑셀 입력 양식: 만들기 → 그대로 읽기(예시 줄은 건너뜀)
    tx = exchange.input_template(["기본사업", "전략개발단"], "인공지능응용연구실", "김연구", "2026-W41")
    sheets = {sh["name"]: sh for sh in xlsx.read(tx)}
    assert list(sheets) == ["작성", "작성법", "과제목록"] and sheets["작성"]["rows"][0][2] == "과제명" and sheets["과제목록"]["rows"][1][0] == "기본사업"
    sx = zipfile.ZipFile(io.BytesIO(tx)).read("xl/worksheets/sheet1.xml").decode()
    assert "<dataValidation" in sx and 'state="frozen"' in sx and "yyyy-mm-dd" in zipfile.ZipFile(io.BytesIO(tx)).read("xl/styles.xml").decode()
    res = exchange.import_file("양식.xlsx", tx, "2026-W41")
    assert res["method"] == "table" and res["rows"] == [], res["rows"][:2]
    try:
        import openpyxl  # 있으면 실제 엑셀 라이브러리로도 열어 본다
        wb = openpyxl.load_workbook(io.BytesIO(tx))
        assert wb.sheetnames == ["작성", "작성법", "과제목록"] and wb["작성"]["C1"].value == "과제명" and wb["작성"].freeze_panes == "A2"
        assert wb["작성"].data_validations.dataValidation and wb["작성"]["E2"].number_format == "yyyy-mm-dd"
        rep_wb = openpyxl.load_workbook(io.BytesIO(app.export({"agg": agg, "format": "xlsx"})[0]))
        assert rep_wb.sheetnames[0] == "보고서" and rep_wb["_data"].sheet_state == "hidden" and rep_wb["보고서"].page_setup.fitToHeight == 1
    except ImportError:
        pass
    # 엑셀 프로그램(openpyxl)으로 열어 채우고 다시 저장한 양식 — 글이 문자 참조(&#…;)·공유 문자열로 바뀌어도 표로 읽혀야 함
    try:
        import openpyxl
        wb2 = openpyxl.load_workbook(io.BytesIO(tx))
        ws2 = wb2["작성"]
        for r, row in enumerate([["인공지능응용연구실", "김연구", "기본사업", "수행", datetime.date(2026, 10, 5), datetime.date(2026, 10, 7), "대리모델 학습", 1, "O"],
                                 ["인공지능응용연구실", "김연구", None, "수행", None, None, "데이터 1,200건 생성", 2, None],
                                 ["인공지능응용연구실", "김연구", "전략개발단", "계획", None, "10/20", "KINS 협의", 1, None]], start=5):
            for c, v in enumerate(row, start=1):
                ws2.cell(r, c, v)
        ob = io.BytesIO()
        wb2.save(ob)
        ro = exchange.import_file("엑셀저장.xlsx", ob.getvalue(), "2026-W41")
        assert ro["method"] == "table" and len(ro["rows"]) == 3, (ro["method"], len(ro["rows"]))
        assert [(x["item"]["project"], x["item"]["period"], x["item"]["depth"], x["item"]["core"]) for x in ro["rows"]] == [
            ("기본사업", "10.5~10.7", 0, True), ("기본사업", "", 1, False), ("전략개발단", "~10.20", 0, False)], ro["rows"]
    except ImportError:
        pass
    assert xlsx._unesc("&#48512;&#49436;") == "부서" and exchange.pptx._unesc("&#44284;&#51228;") == "과제"
    # 지저분한 엑셀: 다른 머리글 이름, 날짜 일련번호·글 날짜, 병합된 과제명, 빈 줄, 쓸모없는 열
    b = xlsx.Book()
    m = b.sheet("Sheet1")
    for c, h in enumerate(["No", "실", "이름", "사업명", "수행/계획", "시작", "마감", "업무 내용", "단계", "중요(핵심)", "과기부", "비공개", "장소", "참석", "아무 열"]):
        m.set(1, c, h)
    m.set(0, 0, "주간보고 (제목 줄)")
    rows_ = [[1, "열수력안전연구실", "이열수", "ATLAS 과제", "실적", 46300, 46303, "ATLAS 시험 1회 수행", 1, "O", "", "", "", "", "x"],
             [2, "열수력안전연구실", "이열수", None, "실적", "'26.10.6.", "10/8", "압력 데이터 정리", 2, "", "", "", "", "", ""],
             [None] * 15,
             [3, "열수력안전연구실", "이열수", None, "예정", None, "2026-10-20", "결과보고서 초안 작성", "", "", "o", "", "", "", ""],
             [4, "열수력안전연구실", "이열수", "", "", "10.13", "", "학회 발표 준비", 1, "", "", "O", "부산 BEXCO", "이열수", ""],
             [5, "열수력안전연구실", "이열수", "기본사업", "모름", "99/99", "", "", 1, "", "", "", "", "", ""]]
    for r, row in enumerate(rows_, start=2):
        for c, v in enumerate(row):
            m.set(r, c, v)
    m.merge(2, 3, 5, 3)  # 과제명 칸 병합(ATLAS 과제 3줄)
    res = exchange.import_file("지저분.xlsx", b.save(), "2026-W41")
    assert res["method"] == "table", res
    rr = res["rows"]
    assert len(rr) == 5, [(x["row"], x["item"]["text"], x["errors"]) for x in rr]
    assert rr[0]["item"]["period"] == "10.5~10.8" and rr[0]["item"]["core"] and rr[0]["item"]["project"] == "ATLAS 과제", rr[0]  # 일련번호 날짜
    assert rr[1]["item"]["depth"] == 1 and rr[1]["item"]["period"] == "10.6~10.8" and rr[1]["item"]["project"] == "ATLAS 과제"  # 글 날짜, 병합 과제명
    assert rr[2]["item"]["kind"] == "plan" and rr[2]["item"]["period"] == "~10.20" and rr[2]["item"]["msit"] and rr[2]["item"]["project"] == "ATLAS 과제"
    assert rr[3]["item"]["kind"] == "plan" and any("추정" in w for w in rr[3]["warnings"]) and rr[3]["item"]["nobbs"] and rr[3]["item"]["place"] == "부산 BEXCO"
    assert any("시작일" in e for e in rr[4]["errors"]) and any("내용 없음" in e for e in rr[4]["errors"]) and any("구분" in e for e in rr[4]["errors"])
    # 형식 섞인 여러 파일 한꺼번에(취합): 데이터형 엑셀(넣은 JSON) + JSON 없는 HWPX + 지저분한 엑셀 + 텍스트
    files = [{"name": "a.xlsx", "b64": base64.b64encode(app.export({"agg": agg, "format": "xlsx-data"})[0]).decode()},
             {"name": "b.hwpx", "b64": base64.b64encode(strip_embed(app.export({"agg": agg, "format": "hwpx"})[0])).decode()},
             {"name": "c.xlsx", "b64": base64.b64encode(b.save()).decode()},
             {"name": "d.txt", "b64": base64.b64encode("이번 주에는 KQZT 시험을 했다".encode()).decode()}]
    res = app.import_files({"files": files, "week": "2026-W41"})
    assert [f["method"] for f in res["files"]] == ["embedded", "form", "table", "memo"], [f["method"] for f in res["files"]]
    assert res["files"][3]["memo"].startswith("이번 주") and any(a["abbr"] == "BEXCO" for a in res["asks"])  # 모르는 약어 → 질의 흐름

    # 12-2) 빈 작성 양식 4종(XLSX·HWPX·DOCX·PPTX) → 프로그램으로 채움 → 올리기 → 기대한 항목
    def zrepl(data, part, fn):
        zin, out = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:
            for i in zin.infolist():
                b_ = zin.read(i)
                z.writestr(i, fn(b_.decode()).encode() if i.filename == part else b_, compress_type=i.compress_type)
        return out.getvalue()
    depts2 = ["가상원자로연구실", "인공지능응용연구실"]
    T = {f: exchange.form_template(f, depts2, "2026-W41", gl, tpl, "선진원자로연구소") for f in ("xlsx", "hwpx", "docx", "pptx")}
    for f, d in T.items():
        r0 = exchange.import_file("빈양식." + f, d, "2026-W41")
        assert r0["method"] in ("table", "form") and r0["rows"] == [], (f, r0["method"], r0["rows"][:1])  # 빈 자리는 항목이 아님
    try:
        import pptx as python_pptx  # 있으면 실제 라이브러리로 열어 본다
        pr = python_pptx.Presentation(io.BytesIO(T["pptx"]))
        assert len(pr.slides) == 1 + 2 + 1 and [s_.name for s_ in pr.slides[1].shapes][:3] == ["실"] + [t_.strip() for t_ in tpl["titles"].values()]
        assert len(python_pptx.Presentation(io.BytesIO(app.export({"agg": agg, "format": "pptx"})[0])).slides) == 1 + 3 + 1
    except ImportError:
        pass
    # XLSX: 셀 채우기(작성 시트 5행부터 = 예시 3줄 아래)
    def xfill(x):
        vals = {"A5": "가상원자로연구실", "B5": "홍길동", "C5": "기본사업", "D5": "수행", "G5": "노심 시험 수행", "H5": "1", "J5": "O",
                "A6": "가상원자로연구실", "B6": "홍길동", "D6": "수행", "G6": "압력 측정", "H6": "2",
                "A7": "인공지능응용연구실", "B7": "김연구", "C7": "전략개발단", "D7": "계획", "G7": "KINS 협의", "H7": "1", "K7": "O"}
        for k, v in vals.items():
            x = re.sub(rf'<c r="{k}"( s="\d+")?/>', lambda m: f'<c r="{k}"{m.group(1) or ""} t="inlineStr"><is><t>{v}</t></is></c>', x, count=1)
        return x.replace('<c r="E5" s="', '<c r="E5" s="').replace('<c r="E5"', '<c r="E5"', 1)
    fx = zrepl(T["xlsx"], "xl/worksheets/sheet1.xml", xfill)
    fx = zrepl(fx, "xl/worksheets/sheet1.xml", lambda x: re.sub(r'<c r="E5"( s="\d+")?/>', lambda m: f'<c r="E5"{m.group(1)}><v>46300</v></c>', x, count=1))
    rx_ = exchange.import_file("채움.xlsx", fx, "2026-W41")["rows"]
    assert [(r_["dept"], r_["item"]["project"], r_["item"]["kind"], r_["item"]["text"], r_["item"]["depth"], r_["item"]["msit"], r_["item"]["nobbs"], r_["item"]["period"]) for r_ in rx_] == [
        ("가상원자로연구실", "기본사업", "done", "노심 시험 수행", 0, True, False, "10.5"), ("가상원자로연구실", "기본사업", "done", "압력 측정", 1, False, False, ""),
        ("인공지능응용연구실", "전략개발단", "plan", "KINS 협의", 0, False, True, "")], rx_
    # HWPX: 첫 실의 수행 칸 '(과제명)' → '(기본사업)', '(기간) 내용' → '(10.5~10.8) 노심 시험 수행' 파랑 글자
    hdr_ = zipfile.ZipFile(io.BytesIO(T["hwpx"])).read("Contents/header.xml").decode()
    bm = re.search(r'<hh:charPr id="(\d+)"[^>]*textColor="#0000FF"', hdr_)
    if not bm:  # 양식에 파랑 글자모양이 없으면(임시 양식) 한글에서 글자색을 바꾼 것처럼 하나 더한다
        last = list(re.finditer(r'<hh:charPr id="(\d+)".*?</hh:charPr>', hdr_, re.S))[-1]
        nid = str(int(last.group(1)) + 1)
        newc = re.sub(r'id="\d+"', f'id="{nid}"', last.group(0), count=1)
        newc = re.sub(r'textColor="[^"]*"', 'textColor="#0000FF"', newc, count=1)
        T["hwpx"] = zrepl(T["hwpx"], "Contents/header.xml", lambda x: re.sub(r'<hh:charProperties itemCnt="(\d+)"', lambda m: f'<hh:charProperties itemCnt="{int(m.group(1)) + 1}"', x.replace("</hh:charProperties>", newc + "</hh:charProperties>", 1), count=1))
        blue = nid
    else:
        blue = bm.group(1)
    def hfill(x):
        x = x.replace("<hp:t>(과제명)</hp:t>", "<hp:t>(기본사업)</hp:t>", 1)
        i = x.index("<hp:t>(기간) 내용</hp:t>")
        j = x.rindex('<hp:run charPrIDRef="', 0, i)
        return x[:j] + f'<hp:run charPrIDRef="{blue}">' + x[x.index(">", j) + 1:i] + "<hp:t>(10.5~10.8) 노심 시험 수행</hp:t>" + x[i + len("<hp:t>(기간) 내용</hp:t>"):]
    rh = exchange.import_file("채움.hwpx", zrepl(T["hwpx"], "Contents/section0.xml", hfill), "2026-W41")
    assert [(r_["dept"], r_["item"]["project"], r_["item"]["period"], r_["item"]["text"], bool(r_["item"].get("msit"))) for r_ in rh["rows"]] == [
        ("가상원자로연구실", "기본사업", "10.5~10.8", "노심 시험 수행", True)], rh["rows"]
    # DOCX: 둘째 실 계획 칸 → '(전략개발단)' / '(~10.20) 보고서 제출' 주황 + 한 줄 더 취소선
    def dfill(x):
        k = [m.start() for m in re.finditer(re.escape(">(과제명)<"), x)][3]  # 순서: 수행(실1, 실2), 계획(실1, 실2)
        x = x[:k] + ">(전략개발단)<" + x[k + len(">(과제명)<"):]
        k2 = [m.start() for m in re.finditer(re.escape(">(기간) 내용<"), x)][3]
        x = x[:k2] + ">(~10.20) 보고서 제출<" + x[k2 + len(">(기간) 내용<"):]
        rs = x.rindex("<w:r>", 0, k2)
        x = x[:rs] + '<w:r><w:rPr><w:color w:val="FF6600"/></w:rPr>' + x[rs + len("<w:r>"):].replace("<w:rPr></w:rPr>", "", 0)
        pe = x.index("</w:p>", k2) + len("</w:p>")
        return x[:pe] + '<w:p><w:pPr><w:ind w:left="374" w:hanging="264"/></w:pPr><w:r><w:t xml:space="preserve">- </w:t></w:r><w:r><w:rPr><w:strike/></w:rPr><w:t xml:space="preserve">내부 협의</w:t></w:r></w:p>' + x[pe:]
    rd_ = exchange.import_file("채움.docx", zrepl(T["docx"], "word/document.xml", dfill), "2026-W41")
    assert [(r_["dept"], r_["item"]["kind"], r_["item"]["project"], r_["item"]["period"], r_["item"]["text"], bool(r_["item"].get("core")), bool(r_["item"].get("nobbs")))
            for r_ in rd_["rows"]] == [("인공지능응용연구실", "plan", "전략개발단", "~10.20", "보고서 제출", True, False),
                                         ("인공지능응용연구실", "plan", "전략개발단", "", "내부 협의", False, True)], rd_["rows"]
    # PPTX: 첫 실 슬라이드 수행 상자 — 수준 0 과제명, 수준 1 항목(주황), 수준 2 하위(취소선)
    def pfill(x):
        x = x.replace("<a:t>(과제명)</a:t>", "<a:t>(기본사업)</a:t>", 1)
        x = x.replace('dirty="0"><a:latin typeface="맑은 고딕"/><a:ea typeface="맑은 고딕"/></a:rPr><a:t>(기간) 내용</a:t>',
                      'dirty="0"><a:solidFill><a:srgbClr val="FF6600"/></a:solidFill><a:latin typeface="맑은 고딕"/><a:ea typeface="맑은 고딕"/></a:rPr><a:t>(10.6) 시험 장치 점검</a:t>', 1)
        i = x.index("(10.6) 시험 장치 점검")
        pe = x.index("</a:p>", i) + len("</a:p>")
        return x[:pe] + '<a:p><a:pPr lvl="2" marL="868680" indent="-228600"><a:buChar char="·"/></a:pPr><a:r><a:rPr lang="ko-KR" strike="sngStrike"/><a:t>계측 채널 교체</a:t></a:r></a:p>' + x[pe:]
    rp = exchange.import_file("채움.pptx", zrepl(T["pptx"], "ppt/slides/slide2.xml", pfill), "2026-W41")
    assert [(r_["dept"], r_["item"]["kind"], r_["item"]["project"], r_["item"]["period"], r_["item"]["text"], r_["item"]["depth"], bool(r_["item"].get("core")), bool(r_["item"].get("nobbs")))
            for r_ in rp["rows"]] == [("가상원자로연구실", "done", "기본사업", "10.6", "시험 장치 점검", 0, True, False),
                                        ("가상원자로연구실", "done", "기본사업", "", "계측 채널 교체", 1, False, True)], rp["rows"]
    # 보고서 PPTX 왕복(넣은 JSON) + JSON 없이 슬라이드 모양만으로
    dp = app.export({"agg": agg, "format": "pptx"})[0]
    assert exchange.import_file("r.pptx", dp, "2026-W41")["method"] == "embedded"
    noj = zipfile.ZipFile(io.BytesIO(dp))
    pbuf = io.BytesIO()
    with zipfile.ZipFile(pbuf, "w") as z:
        for i in noj.infolist():
            if "customXml" not in i.filename:
                z.writestr(i, noj.read(i))
    rq = exchange.import_file("r.pptx", pbuf.getvalue(), "2026-W41")
    gotp = {}
    for r_ in rq["rows"]:
        gotp.setdefault(r_["dept"], []).append(app.norm_item(r_["item"]))
    cmpp = lambda its: norm([{k: v for k, v in i.items() if k not in ("place", "people")} | {"text": i["text"] + rules.ext_suffix(i)} for i in its])
    A2, B2 = {k: cmpp(v) for k, v in gotp.items()}, {r_["label"]: cmpp(r_["items"]) for r_ in agg["rows"] if r_["items"]}
    assert rq["method"] == "form" and A2 == B2 and rq["org"] == "가상원자력연구소", (rq["method"], rq["org"], [(k, set(A2.get(k, [])) ^ set(B2.get(k, []))) for k in set(A2) | set(B2) if A2.get(k) != B2.get(k)])
    # 형식 섞어 한꺼번에(취합): 채운 XLSX + HWPX + DOCX + PPTX
    mixed = app.import_files({"files": [{"name": n_, "b64": base64.b64encode(d_).decode()} for n_, d_ in (("a.xlsx", fx), ("b.hwpx", zrepl(T["hwpx"], "Contents/section0.xml", hfill)),
                                         ("c.docx", zrepl(T["docx"], "word/document.xml", dfill)), ("d.pptx", zrepl(T["pptx"], "ppt/slides/slide2.xml", pfill)))], "week": "2026-W41"})
    assert [len(f_["rows"]) for f_ in mixed["files"]] == [3, 1, 2, 2]

    # 12-3) 과제 등록부(양식 만들기): 추가·이름 바꿈·삭제·되살림·순서·구성원·지난 제출에서 가져오기 → 등록부로 채운 작성 양식 4종 → 채움 → 올리기
    O, D = "가상원자력연구소", "인공지능응용연구실"
    v = app.registry_op({"op": "add", "org": O, "dept": D, "name": "기본사업", "editor": "관리자"})
    app.registry_op({"op": "add", "org": O, "dept": D, "name": "(i-SMR 과제)", "full": "혁신형 SMR 노심 과제", "period": "10/5~10/9", "editor": "관리자"})
    app.registry_op({"op": "add", "org": O, "dept": D, "name": "전략개발단", "person": "박서준", "editor": "관리자"})
    u = app.registry_view(O, D)["units"][0]
    ids = {p_["name"]: p_["id"] for p_ in u["projects"]}
    assert list(ids) == ["기본사업", "i-SMR 과제", "전략개발단"] and u["projects"][1]["period"] == "10.5~10.9" and u["projects"][0]["editor"] == "관리자"
    app.registry_op({"op": "update", "org": O, "dept": D, "id": ids["i-SMR 과제"], "name": "i-SMR 노심 과제", "editor": "김연구"})
    app.registry_op({"op": "delete", "org": O, "dept": D, "id": ids["기본사업"], "editor": "김연구"})
    u = app.registry_view(O, D)["units"][0]
    assert [p_["name"] for p_ in u["projects"]] == ["i-SMR 노심 과제", "전략개발단"] and [p_["name"] for p_ in u["deleted"]] == ["기본사업"]
    hist_ = next(p_ for p_ in u["projects"] if p_["id"] == ids["i-SMR 과제"])["history"]
    assert hist_[-1]["action"] == "이름 바꿈" and hist_[-1]["before"]["name"] == "i-SMR 과제" and hist_[-1]["editor"] == "김연구"
    app.registry_op({"op": "restore", "org": O, "dept": D, "id": ids["기본사업"], "editor": "김연구"})
    app.registry_op({"op": "reorder", "org": O, "dept": D, "ids": [ids["전략개발단"], ids["i-SMR 과제"], ids["기본사업"]]})
    app.registry_op({"op": "members", "org": O, "dept": D, "members": "김연구, 박서준"})
    u = app.registry_view(O, D)["units"][0]
    assert [p_["name"] for p_ in u["projects"]] == ["전략개발단", "i-SMR 노심 과제", "기본사업"] and u["members"] == ["김연구", "박서준"]
    assert app.known_projects({"org": O, "dept": D, "name": "김연구"}, "2026-W41")[:2] == ["i-SMR 노심 과제", "기본사업"]  # 김연구에게는 박서준 과제 빼고, 등록 순서
    vv = app.registry_op({"op": "copy_prev", "org": O, "dept": D, "week": "2026-W42", "editor": "관리자"})  # 지난 제출에서 쓴 과제명 가져오기
    assert "AI융합사업" in [p_["name"] for p_ in vv["units"][[x["dept"] for x in vv["units"]].index(D)]["projects"]]
    app.registry_op({"op": "delete", "org": O, "dept": D, "id": next(p_["id"] for p_ in app.registry_view(O, D)["units"][0]["projects"] if p_["name"] == "AI융합사업")})
    # 등록부로 채운 양식: 활성 과제마다 '∙ (과제명)' + 빈 '- (기간) 내용', 사람별 블록
    prev_ = app.registry_template({"org": O, "dept": D, "week": "2026-W41", "per_person": True, "preview": True})
    assert prev_["rows"] == 2 and prev_["html"].count("(전략개발단)") == 2 and prev_["html"].count("(i-SMR 노심 과제)") == 4  # 박서준만 전략개발단
    assert "i-SMR 과제)" not in prev_["html"]  # 바꾼 이름이 새 양식에 반영
    R = {f: app.registry_template({"org": O, "dept": D, "week": "2026-W41", "fmt": f, "prefill": True})[0] for f in ("xlsx", "hwpx", "docx", "pptx")}
    for f, d_ in R.items():
        assert exchange.import_file("등록양식." + f, d_, "2026-W41")["rows"] == [], f  # 빈 자리는 항목 아님
    pre_rows = [r_ for r_ in xlsx.read(R["xlsx"])[0]["rows"][4:10] if r_ and r_[2]]
    assert [(r_[1], r_[2]) for r_ in pre_rows] == [("김연구", "i-SMR 노심 과제"), ("김연구", "기본사업"), ("박서준", "전략개발단"), ("박서준", "i-SMR 노심 과제"),
                                                  ("박서준", "기본사업")] and xlsx.read(R["xlsx"])[2]["rows"][1][0] == "전략개발단"  # 과제목록 시트도 등록부
    fill1 = lambda x, a, b: x.replace(a, b, 1)
    hh = zrepl(R["hwpx"], "Contents/section0.xml", lambda x: fill1(x, "<hp:t>(10.5~10.9) </hp:t>", "<hp:t>(10.5~10.9) </hp:t>").replace("<hp:t>(기간) 내용</hp:t>", "<hp:t>(10.6) 노심 해석 착수</hp:t>", 1))
    dd = zrepl(R["docx"], "word/document.xml", lambda x: x.replace(">(기간) 내용<", ">(10.7) 전략 회의<", 1))
    pp = zrepl(R["pptx"], "ppt/slides/slide2.xml", lambda x: x.replace("<a:t>내용</a:t>", "<a:t>노심 코드 정리</a:t>", 1))
    xx = zrepl(R["xlsx"], "xl/worksheets/sheet1.xml", lambda x: re.sub(r'<c r="G6"( s="\d+")?/>', lambda m: f'<c r="G6"{m.group(1)} t="inlineStr"><is><t>노심 해석 보고</t></is></c>', x)
               .replace('<c r="D6"', '<c r="D6" t="inlineStr"', 1).replace('<c r="D6" t="inlineStr" s="', '<c r="D6" t="inlineStr" s="', 1))
    xx = zrepl(xx, "xl/worksheets/sheet1.xml", lambda x: re.sub(r'<c r="D6" t="inlineStr"( s="\d+")?/>', lambda m: f'<c r="D6"{m.group(1)} t="inlineStr"><is><t>수행</t></is></c>', x))
    got_ = {f: [(r_["item"]["project"], r_["item"]["period"], r_["item"]["text"]) for r_ in exchange.import_file("채움." + f, d_, "2026-W41")["rows"]]
            for f, d_ in (("hwpx", hh), ("docx", dd), ("pptx", pp), ("xlsx", xx))}
    assert got_["hwpx"] == [("전략개발단", "", "노심 해석 착수")] or got_["hwpx"] == [("전략개발단", "10.6", "노심 해석 착수")], got_
    assert got_["docx"] == [("전략개발단", "10.7", "전략 회의")] and got_["pptx"] == [("i-SMR 노심 과제", "10.5~10.9", "노심 코드 정리")], got_
    assert got_["xlsx"] == [("기본사업", "", "노심 해석 보고")], got_
    # 가져올 때 등록 과제와 비슷한 이름 → 제안(바꾸지 않음)
    nb = xlsx.Book()
    ns = nb.sheet("작성")
    for c, h in enumerate(["부서", "작성자", "과제명", "구분", "내용"]):
        ns.set(0, c, h)
    for c, v_ in enumerate([D, "김연구", "i-SMR노심과제", "수행", "노심 시험"]):
        ns.set(1, c, v_)
    ri = app.import_files({"files": [{"name": "n.xlsx", "b64": base64.b64encode(nb.save()).decode()}], "week": "2026-W41"})["files"][0]["rows"][0]
    assert ri["item"]["project"] == "i-SMR노심과제" and ri["suggest_project"] == "i-SMR 노심 과제" and any("제안만" in w for w in ri["warnings"])
    # 취합 순서: 등록부 순서대로
    od = app.agg_doc({"week": "2026-W41", "org": O, "rows": [{"label": D, "items": [
        {"id": "a", "text": "가", "kind": "done", "project": "기본사업"}, {"id": "b", "text": "나", "kind": "done", "project": "전략개발단"}]}]})
    oh = render.to_html(od, gl, tpl)
    assert oh.index("(전략개발단)") < oh.index("(기본사업)")

    # 12-4) 과제 별칭·합치기 제안·참석자 표기·한글 풀이만 있는 약어·약어 목록 하나
    OD = ("가상원자력연구소", "인공지능응용연구실")
    app.registry_op({"op": "alias", "org": OD[0], "dept": OD[1], "name": "주간보고서 자동화", "alias": "주간보고서 봇", "editor": "김연구"})
    assert app.REG.canonical("주간보고서봇", *OD) == ("주간보고서 자동화", "별칭") and app.REG.canonical("주간보고서 자동화", *OD)[0] == "주간보고서 자동화"
    assert app.REG.canonical("주간보고서자동화x", *OD) == (None, "")  # 등록 이름과 표기만 비슷한 것은 바꾸지 않음(제안 경로)
    its_ = [app.norm_item({"text": "초안 생성 기능 구현", "project": "주간보고서 자동화"}), app.norm_item({"text": "자동 분류 보완", "project": "주간보고서 봇"}),
            app.norm_item({"text": "추출 규칙 수정", "project": "주간보고서 봇", "depth": 1})]
    log_ = app.apply_aliases(its_, *OD)
    assert log_ == ["과제명 '주간보고서 봇' → '주간보고서 자동화'(별칭)"] and {i["project"] for i in its_} == {"주간보고서 자동화"}
    # 합치기 제안: 규칙(LLM 없이) — 4글자 이상 핵심 낱말 공유는 제안, '자동화'(3글자)만 같은 다른 과제는 제안 안 함, 따로 두기는 기억
    g_ = [app.norm_item({"text": t, "project": p_}) for t, p_ in (("a", "보고서봇 개발"), ("b", "보고서봇"), ("c", "문서 자동화"), ("d", "영상 자동화"))]
    sg = app.suggest_merges(g_, "", "테스트실", use_llm=False)
    assert [(m["a"], m["into"]) for m in sg] == [("보고서봇", "보고서봇 개발")] or [(m["a"], m["into"]) for m in sg] == [("보고서봇 개발", "보고서봇")], sg
    app.dismiss_merge({"dept": "테스트실", "a": "보고서봇", "b": "보고서봇 개발"})
    assert app.suggest_merges(g_, "", "테스트실", use_llm=False) == []
    n0 = len(CALLS)
    app.suggest_merges(g_ + [app.norm_item({"text": "e", "project": "내부망 AI"}), app.norm_item({"text": "f", "project": "내부망 AI 업무지원"})], "", "테스트실")
    assert len(CALLS) == n0 + 1  # 후보가 있으면 LLM 한 번(예/아니오)
    # 참석자 표기: 외부활동만(기본) — 내부 동료·내부 장소·전화는 안 붙임, 문장에 있는 이름은 다시 안 붙임
    E = lambda **k: rules.ext_suffix(dict({"text": "", "place": "", "people": "", "ext": False}, **k))
    assert E(text="PC 로컬 실행 테스트", people="김선임") == "" and E(text="화면 의견 정리", place="본관 회의실", people="김선임, 박책임", ext=True) == ""
    assert E(text="후보 기준 협의(전화)", people="서박사", ext=True) == "" and E(text="데모", people="팀장") == ""
    assert E(text="IAEA 회의 발표", place="오스트리아 빈", people="김연구", ext=True) == " (@오스트리아 빈, 김연구)"
    assert E(text="김연구가 KINS 협의 참석", place="KINS 대전", people="김연구", ext=True) == " (@KINS 대전)"
    assert rules.ext_suffix({"text": "데모", "people": "팀장"}, "always") == " (팀장)" and rules.ext_suffix({"text": "x", "place": "빈", "ext": True}, "never") == ""
    # 한글 풀이만 있는 약어(이상탐지(VAD), 오탐(FP))도 ※ 약어에 영문 전체 이름으로, 영문 이름이 문장에 있으면 생략
    o_, n_, d_, u_, _ = rules.expand_items(["영상 이상탐지(VAD)에서 오탐(FP) 확인", "Video Anomaly Detection(VAD) 재검토"], gl, mode="note")
    assert [a for a, _ in n_][:2] == ["VAD", "FP"] and "Video Anomaly Detection" in dict(n_)["VAD"] and "False Positive" in dict(n_)["FP"], n_
    o_, n_, d_, u_, _ = rules.expand_items(["Video Anomaly Detection(VAD) 재검토"], gl, mode="note")
    assert n_ == []
    # 설명을 켜도(수준 3 + 설명 따로) 약어 목록은 하나 — '약어: 이름 — 설명'
    o_, n_, d_, u_, _ = rules.expand_items(["RAG 구축"], gl, mode="note", desc_block=True)
    assert d_ == [] and " — " in dict(n_)["RAG"]
    ddoc = app.personal_doc({"week": "2026-W41", "dept": "d", "name": "n", "items": [app.norm_item({"text": "RAG 구축", "project": "p"})]}, {"desc_block": True})
    tt = render.to_text(ddoc, gl, tpl)
    assert tt.count("RAG:") == 1 and "※ 용어 설명" not in tt
    # 등록부가 있으면 항목 정리 프롬프트는 등록 과제 목록(닫힌 목록 + 별칭)으로
    app.run_itemize({"profile": {"name": "김연구", "org": OD[0], "dept": OD[1]}, "week": "2026-10-07", "memo": "- 시험 수행", "use_prev": False}, EMIT, "fake")
    it_call = [u for sy, u in CALLS if sy.startswith("너는 한국원자력연구원(KAERI)")][-1]
    assert "[등록 과제 — project 는 반드시 이 중에서 고른다" in it_call and "- 주간보고서 자동화 (별칭: 주간보고서 봇)" in it_call

    # 12-5) 과기정통부·핵심은 명시 표시만(언급만이면 질문), 장소·상대·참석자 나누기, 작성자 포함, 온라인, 연구원 안 장소는 문장에
    #       (실제 gemma4:31b 가 이 메모에서 낸 출력을 그대로 흉내 — 결정론 부분 검사)
    MIX_MEMO = """이번 주
- [i-SMR 과제] 과기정통부 원자력정책과에 3분기 실적 자료 제출 (과기정통부 보고)
- [i-SMR 과제] 세종 정부청사에서 과기정통부 김사무관, 우리 실 박책임과 함께 사업 진도 점검 회의
- [AI 응용과제] 대전 KAIST에서 이교수님과 공동연구 협의, 김선임 동행
- [AI 응용과제] 본관 대회의실에서 원장님 대상 AI 성과 보고 ★
- 과기정통부 요청으로 SMR 안전성 자료 검토 (내부, 박책임)
- [기본사업] IAEA 기술회의 발표 (@오스트리아 빈, 홍길동)
- [기본사업] KINS 담당자와 화상회의로 인허가 일정 협의
다음 2주
- [i-SMR 과제] 과기정통부 장관 현장 방문 대응 준비 (10.20, 연구원 본관)
- [AI 응용과제] 한국원자력학회 추계학술대회 발표 (@창원 CECO, 홍길동·이연구)"""
    MIX_RAW = '{"items": [{"project": "i-SMR 과제", "period": "", "text": "과기정통부 원자력정책과 3분기 실적 자료 제출", "kind": "done", "cat": "perf", "ext": false, "place": "", "people": "", "core": false, "msit": true, "nobbs": false}, {"project": "i-SMR 과제", "period": "", "text": "과기정통부 사업 진도 점검 회의", "kind": "done", "cat": "goal", "ext": true, "place": "세종 정부청사", "people": "과기정통부 김사무관, 박책임", "core": false, "msit": false, "nobbs": false}, {"project": "AI 응용과제", "period": "", "text": "이교수 공동연구 협의", "kind": "done", "cat": "goal", "ext": true, "place": "대전 KAIST", "people": "이교수, 김선임", "core": false, "msit": false, "nobbs": false}, {"project": "AI 응용과제", "period": "", "text": "원장 대상 AI 성과 보고", "kind": "done", "cat": "perf", "ext": false, "place": "", "people": "", "core": true, "msit": false, "nobbs": false}, {"project": "", "period": "", "text": "SMR 안전성 자료 검토", "kind": "done", "cat": "etc", "ext": false, "place": "", "people": "", "core": false, "msit": true, "nobbs": false}, {"project": "기본사업", "period": "", "text": "IAEA 기술회의 발표", "kind": "done", "cat": "perf", "ext": true, "place": "오스트리아 빈", "people": "홍길동", "core": false, "msit": false, "nobbs": false}, {"project": "기본사업", "period": "", "text": "KINS 담당자 인허가 일정 협의", "kind": "done", "cat": "goal", "ext": true, "place": "온라인", "people": "KINS 담당자", "core": false, "msit": false, "nobbs": false}, {"project": "i-SMR 과제", "period": "10.20", "text": "과기정통부 장관 현장 방문 대응 준비", "kind": "plan", "cat": "goal", "ext": false, "place": "", "people": "", "core": false, "msit": false, "nobbs": false}, {"project": "AI 응용과제", "period": "", "text": "한국원자력학회 추계학술대회 발표", "kind": "plan", "cat": "perf", "ext": true, "place": "창원 CECO", "people": "홍길동, 이연구", "core": false, "msit": false, "nobbs": false}], "carry": [], "remarks": []}'
    app.llm = lambda system, user, *a, **k: MIX_RAW if "주간보고 메모를" in system else fake(system, user, *a, **k)
    try:
        mx = app.run_itemize({"profile": {"name": "홍길동", "org": "가상원자력연구소", "dept": "혼합테스트실"}, "week": "2026-10-07", "memo": MIX_MEMO, "use_prev": False}, EMIT, "fake")
    finally:
        app.llm = fake
    MX = {i["text"]: i for i in mx["items"]}
    sfx = lambda i: i["text"] + rules.ext_suffix(i, "external", True, "홍길동")
    a1 = MX["과기정통부 원자력정책과 3분기 실적 자료 제출"]
    assert a1["msit"] and not a1.get("ask")  # '(과기정통부 보고)' 명시 → 파랑
    a5 = MX["SMR 안전성 자료 검토"]
    assert not a5["msit"] and a5["ask"] == ["msit"] and not a5["people"] and not a5["place"]  # 과기정통부 '요청'은 보고 사항 아님 → 질문
    a2 = next(i for i in mx["items"] if "사업 진도 점검 회의" in i["text"])
    assert a2["text"] == "과기정통부 김사무관과 사업 진도 점검 회의" and a2["party"] == "과기정통부 김사무관" and a2["people"] == "박책임", a2
    assert sfx(a2) == "과기정통부 김사무관과 사업 진도 점검 회의 (@세종 정부청사, 홍길동, 박책임)" and "party" in a2["guess"]
    a3 = next(i for i in mx["items"] if "공동연구 협의" in i["text"])
    assert a3["party"] == "이교수" and a3["people"] == "김선임" and a3["went"] is True and sfx(a3).endswith("(@대전 KAIST, 홍길동, 김선임)"), sfx(a3)
    a4 = next(i for i in mx["items"] if "성과 보고" in i["text"])
    assert a4["core"] and a4["text"].startswith("본관 대회의실에서 ") and sfx(a4) == a4["text"]  # ★ → 핵심, 내부 장소는 문장에
    a6 = MX["IAEA 기술회의 발표"]
    assert sfx(a6) == "IAEA 기술회의 발표 (@오스트리아 빈, 홍길동)"  # 작성자가 이미 있으면 두 번 안 씀
    a7 = next(i for i in mx["items"] if "인허가 일정 협의" in i["text"])
    assert a7["place"] == "온라인" and a7["party"] == "KINS 담당자" and not a7["people"] and sfx(a7).endswith("(@온라인, 홍길동)") and "place" not in (a7.get("guess") or [])
    assert rules.ext_suffix(a7, "external", True, "홍길동", "omit") == " (홍길동)" and rules.ext_suffix(a7, "external", False, "홍길동") == " (@온라인)"
    a8 = next(i for i in mx["items"] if "장관 현장 방문" in i["text"])
    assert a8["text"] == "연구원 본관 과기정통부 장관 현장 방문 대응 준비" and sfx(a8) == a8["text"] and a8["ask"] == ["msit"] and not a8["msit"]
    a9 = next(i for i in mx["items"] if "추계학술대회" in i["text"])
    assert sfx(a9).endswith("(@창원 CECO, 홍길동, 이연구)")
    assert not [n for n in mx["notes"] if "장소" in n]  # 온라인을 '원문에 없는 장소'로 지우지 않음
    xd = app.personal_doc({"week": "2026-W41", "dept": "혼합테스트실", "name": "홍길동", "items": mx["items"]}, {})
    xc = rules.check_doc(xd, gl)
    assert sorted((c["flag"], c["item"]) for c in xc if c["kind"] == "flag") == sorted([("msit", a2["id"]), ("msit", a5["id"]), ("msit", a8["id"])])
    assert not [c for c in xc if c["kind"] == "ext"], [c["msg"] for c in xc if c["kind"] == "ext"]  # 온라인·작성자 포함 → 장소·참석자 없음 경고 없음
    xt = render.to_text(xd, gl, tpl)
    assert "(@세종 정부청사, 홍길동, 박책임)" in xt and "과기정통부 김사무관과" in xt
    assert "(@대전 KAIST, 김선임)" in render.to_text(app.personal_doc({"week": "2026-W41", "dept": "d", "name": "홍길동", "items": [a3]}, {"writer": "no"}), gl, tpl)  # 작성자 자동 포함 끔
    # 다른 사람이 주어이면(대리 참석) 작성자를 넣지 않음, 메모에 (@장소, 참석자) 를 적었으면 그대로
    w1 = app.norm_item({"text": "학회 발표", "ext": True, "place": "부산 BEXCO", "people": "김선임"})
    rules.apply_source(w1, "김선임 학회 발표 대리 참석", "홍길동")
    w2 = app.norm_item({"text": "워크숍 참석", "ext": True, "place": "서울", "people": "박책임"})
    rules.apply_source(w2, "워크숍 참석 (@서울, 박책임)", "홍길동")
    assert not w1.get("went") and not w2.get("went") and "people" not in (w2.get("guess") or [])
    # LLM 이 상대로 둔 우리 쪽 사람(서박사)은 참석자로 → 전화 협의는 외부활동 아님 / 장소 없는 외부 협의엔 작성자만 붙이지 않음
    w3 = app.norm_item({"text": "서박사와 전화로 후보 선정 기준 협의", "ext": True, "party": "서박사"})
    rules.apply_source(w3, "서박사님과 전화로 후보 선정 기준 및 추가 계산 범위 협의", "홍길동")
    w4 = app.norm_item({"text": "정읍 연구팀과 기준 재협의", "ext": True, "party": "정읍 연구팀"})
    rules.apply_source(w4, "[AI Scientist] 정읍 연구팀과 상위 후보물질 선정 기준 재협의", "홍길동")
    assert w3["people"] == "서박사" and not w3["party"] and rules.ext_suffix(w3, "external", True, "홍길동") == "" and w4["party"] == "정읍 연구팀" and not w4.get("went")
    # 사람 칸 나누기·문장에 상대 넣기
    assert rules.split_people("과기정통부 김사무관, 우리 실 박책임, KINS 담당자, 이교수님, 홍길동, 서박사", "홍길동") == (["박책임", "홍길동", "서박사"], ["과기정통부 김사무관", "KINS 담당자", "이교수님"])
    assert rules.party_in_text("인허가 협의", "KINS") == "KINS와 인허가 협의" and rules.party_in_text("KINS 담당자 협의", "KINS 담당자") == "KINS 담당자 협의"
    # 집계(취합): 상대·작성자도 합쳐짐 / 엑셀 '상대' 칸 왕복
    xa = app.merge_dept("d", [dict(a2, sid="s1", who="홍길동"), dict(a2, sid="s2", who="박책임", went=False, text="과기정통부 김사무관과 사업 진도 점검 회의 참석")], "fake", EMIT, 2026)[0]
    assert all(i.get("party") == "과기정통부 김사무관" for i in xa) and [rules.went_names(i) for i in xa] == [["홍길동"], []], xa
    import xlsx as xl
    for head, row in ((["부서", "작성자", "과제명", "구분", "내용", "장소", "참석자", "상대 기관·인물"], ["d", "홍길동", "p", "수행", "진도 점검 회의", "세종", "박책임", "과기정통부 김사무관"]),
                      (["부서", "작성자", "과제명", "구분", "내용", "장소", "참석자"], ["d", "홍길동", "p", "수행", "진도 점검 회의", "세종", "과기정통부 김사무관, 박책임"])):  # 옛 양식(상대 칸 없음)
        bk = xl.Book(); sh = bk.sheet("작성")
        for c_, v_ in enumerate(head):
            sh.set(0, c_, v_)
        for c_, v_ in enumerate(row):
            sh.set(1, c_, v_)
        ri = app.import_files({"files": [{"name": "상대.xlsx", "b64": base64.b64encode(bk.save()).decode()}], "week": "2026-W41"})["files"][0]["rows"][0]["item"]
        assert ri["party"] == "과기정통부 김사무관" and ri["people"] == "박책임" and ri["text"] == "과기정통부 김사무관과 진도 점검 회의", ri
    assert "party" in [k for k, _, _ in exchange.FIELDS] and exchange.map_headers(["상대 기관·인물", "외부활동 참석자(우리 연구원)"])

    # 13) 내용 수집 보조 (다른 도구 기록, 읽기 전용)
    md = os.path.join(TMP, "meeting-local", "2026-10-06-ab12")
    os.makedirs(md)
    open(os.path.join(md, "state.json"), "w").write(json.dumps({"ts": "2026-10-06T10:00:00"}))
    open(os.path.join(md, "minutes.md"), "w").write("# 과제 주간 회의\n- 시험 일정 확정")
    hd = os.path.join(TMP, "mail-local", "history")
    os.makedirs(hd)
    open(os.path.join(hd, "x.json"), "w").write(json.dumps({"id": "x", "kind": "write", "title": "자료 요청", "ts": "2026-10-07T09:00:00", "versions": [{"body": "본문"}]}))
    s = app.sources("2026-W41")
    assert s["meetings"][0]["title"] == "과제 주간 회의" and s["mails"][0]["title"] == "자료 요청"
    assert not app.sources("2026-W40")["meetings"]

    # 14) HTTP (SSE)
    srv = app.ThreadingHTTPServer(("127.0.0.1", 0), app.H)
    port = srv.server_address[1]
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{port}"
    page = urllib.request.urlopen(base + "/").read().decode()
    assert "data-sig" in page and "주간보고" in page
    m = json.load(urllib.request.urlopen(base + "/api/meta"))
    assert m["week"]["key"] and m["glossary"]["public"] > 7000 and m["has_sources"]
    post = lambda p, b: urllib.request.urlopen(urllib.request.Request(base + p, json.dumps(b).encode(), {"Content-Type": "application/json"}))
    pv = json.load(post("/api/preview", {"agg": agg}))
    assert "wk-table" in pv["html"] and isinstance(pv["checks"], list)
    resp = post("/api/export", {"agg": agg, "format": "hwpx"})
    assert resp.headers["Content-Disposition"].startswith("attachment") and resp.read()[:2] == b"PK"
    assert int(resp.headers["X-Unresolved"]) >= 1  # NURETH — 내보내기 경고용
    try:
        post("/api/submit", {"profile": {"name": "h", "dept": "d"}, "items": [{"text": "XQWZ 시험", "kind": "done"}]})
        raise AssertionError("409 아님")
    except urllib.error.HTTPError as e:
        assert e.code == 409 and json.load(e)["unresolved"][0]["abbr"] == "XQWZ"
    sse = post("/api/itemize", {"profile": P1, "week": "2026-10-07", "memo": memo}).read().decode()
    done = [json.loads(l[6:]) for l in sse.split("\n\n") if l.startswith("data: ")]
    assert any("token" in e for e in done) and "done" in done[-1] and len(done[-1]["done"]["items"]) == 11
    g = json.load(urllib.request.urlopen(base + "/api/glossary?q=LOCA"))
    assert any(r["abbr"] == "LOCA" for r in g["rows"])
    srv.shutdown()
    print("selftest OK — 과제 등록부(추가·이름 바꿈·삭제·되살림·순서·구성원·지난 제출)·등록부 양식 4종 채움→올리기·비슷한 과제명 제안·등록 순서·작성 양식 4종(XLSX·HWPX·DOCX·PPTX) 채움→올리기·PPTX 내보내기·공식 양식·과제명 묶음·가져오기(넣은 JSON 왕복·엑셀 입력 양식·손으로 쓴 HWPX/DOCX/엑셀·지저분한 엑셀·여러 형식)·엑셀 내보내기·풀이 줄(굵게)·표 2개(수행/계획)·하위 항목 들여쓰기·주차·약어집(사내>시드>공개·동음이의)·풀이 수준·약어 질의·제출 차단·외부활동(장소·상대·참석자·작성자 포함·온라인)·과기정통부·핵심 명시 표시(언급만이면 질문)·날짜·분량·원문 대조·항목화·지난 계획·제출·미제출·취합·복원·압축·HWPX·DOCX·BBS·수집·HTTP"
          + ("" if kc is None else " · kordoc validate"))
finally:
    shutil.rmtree(TMP, ignore_errors=True)
